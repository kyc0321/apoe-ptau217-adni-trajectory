import re
from docx import Document

PATH = "/Users/kuan-yuchu/Desktop/LOAD-BC/submit BC 5:31:26/sumit 6:3 26/Main_Text.docx"
d = Document(PATH)

def replace_in_para(anchor_substr, old, new):
    p = next(pp for pp in d.paragraphs if anchor_substr in pp.text)
    assert old in p.text, f"'{old}' not found in paragraph anchored by '{anchor_substr[:40]}'"
    full = p.text
    new_full = full.replace(old, new, 1)
    for r in list(p.runs):
        r.text = ""
    p.runs[0].text = new_full
    return p

# ---------------------------------------------------------------------------
# Insert missing in-text citation brackets for references that exist in the
# list but were never actually cited anywhere in the prose (found during a
# citation-integrity audit). Uses each reference's CURRENT number; final
# renumbering (after removing the one genuinely-unused reference, #34) is
# done in a separate consolidated pass.
# ---------------------------------------------------------------------------

# 22 (Yamazaki), 23 (Serrano-Pozo): APOE pathobiology reviews
replace_in_para(
    "APOE ε4 is the strongest common genetic risk factor",
    "affects amyloid clearance, lipid metabolism, microglial function, and reactive astrocytosis [15,16].",
    "affects amyloid clearance, lipid metabolism, microglial function, and reactive astrocytosis [15,16,22,23]."
)

# 24 (Fortea): APOE4 homozygosity as distinct genetic form -- cite alongside
# the mechanistic APOE statement in Discussion §4.3
replace_in_para(
    "APOE ε4 is thought to disrupt amyloid clearance and lipid homeostasis",
    "with downstream consequences for tau pathology [14–16].",
    "with downstream consequences for tau pathology [14–16,24]."
)

# 25 (Liddelow), 26 (Habib): reactive/disease-associated astrocyte biology
replace_in_para(
    "Plasma glial fibrillary acidic protein (GFAP), measured by single-molecule array, marks reactive astrocytosis",
    "marks reactive astrocytosis, rises early on the AD continuum,",
    "marks reactive astrocytosis [25,26], rises early on the AD continuum,"
)

# 35 (Mielke): plasma p-tau181/217 performance across platforms
replace_in_para(
    "Plasma phosphorylated tau-217 (p-tau217), now available on multiple platforms",
    "predicting cognitive decline with discrimination approaching that of CSF p-tau [6–8].",
    "predicting cognitive decline with discrimination approaching that of CSF p-tau [6–8,35]."
)

# 27 (Petersen ADNI characterization), 28 (Weiner ADNI3): cohort description
replace_in_para(
    "Data were obtained from the ADNI database",
    "Data were obtained from the ADNI database (adni.loni.usc.edu) under an approved Data Use Agreement [20], 2026-05-24 release.",
    "Data were obtained from the ADNI database (adni.loni.usc.edu) under an approved Data Use Agreement [20], 2026-05-24 release [27,28]."
)

# 29 (Trojanowski, ADNI biomarker core), 30 (Rissin, Simoa technology)
replace_in_para(
    "Plasma GFAP was measured by single-molecule array (Simoa; Quanterix Corp.)",
    "Plasma GFAP was measured by single-molecule array (Simoa; Quanterix Corp.) on the UPENN Plasma Biomarker panel (column GFAP_Q) [21].",
    "Plasma GFAP was measured by single-molecule array (Simoa; Quanterix Corp.) [30] on the UPENN Plasma Biomarker panel (column GFAP_Q) [21,29]."
)

# 31 (Mohs, ADAS-Cog), 32 (Folstein, MMSE), 33 (Fischl, FreeSurfer)
# NOTE: same paragraph as the Simoa/29/30 edit above (§2.2 combines assay and
# outcome-measure descriptions in one paragraph) -- anchor on that paragraph.
replace_in_para(
    "Plasma GFAP was measured by single-molecule array (Simoa; Quanterix Corp.)",
    "Cognitive outcomes were ADAS-Cog 13 total (column TOTAL13; primary; higher = worse) and MMSE total (column MMSCORE; secondary). Hippocampal volume was the sum of Left.Hippocampus_combat + Right.Hippocampus_combat from the ADSP-harmonized FreeSurfer pipeline,",
    "Cognitive outcomes were ADAS-Cog 13 total (column TOTAL13; primary; higher = worse) [31] and MMSE total (column MMSCORE; secondary) [32]. Hippocampal volume was the sum of Left.Hippocampus_combat + Right.Hippocampus_combat from the ADSP-harmonized FreeSurfer pipeline [33],"
)

# 36 (Inker, CKD-EPI 2021 race-free equation)
replace_in_para(
    "estimated glomerular filtration rate (CKD-EPI 2021 race-free equation",
    "estimated glomerular filtration rate (CKD-EPI 2021 race-free equation; computed from ADNI LABDATA serum creatinine, column RCT392).",
    "estimated glomerular filtration rate (CKD-EPI 2021 race-free equation [36]; computed from ADNI LABDATA serum creatinine, column RCT392)."
)

# 37 (Janelidze, kidney dysfunction / p-tau ratio)
replace_in_para(
    "Plasma GFAP and p-tau217 concentrations are influenced by renal function",
    "Plasma GFAP and p-tau217 concentrations are influenced by renal function (eGFR), body-mass index (BMI), and age [17,18];",
    "Plasma GFAP and p-tau217 concentrations are influenced by renal function (eGFR) [37], body-mass index (BMI), and age [17,18];"
)

# 38 (Proust-Lima, lcmm)
replace_in_para(
    "Latent class growth mixed models (lcmm R package",
    "Latent class growth mixed models (lcmm R package v1.9.2) were fitted",
    "Latent class growth mixed models (lcmm R package [38] v1.9.2) were fitted"
)

# 39 (Bates, lme4), 40 (Bürkner, brms), 42 (Rizopoulos, JMbayes2)
replace_in_para(
    "Analyses were performed in R 4.5.3 with lme4",
    "Analyses were performed in R 4.5.3 with lme4 1.1-37, lmerTest 3.1-3, lcmm 1.9.2, brms 2.21.0, JMbayes2 0.5-1 (attempted then abandoned — see SAP §10 amendment 6), and nlme 3.1-164.",
    "Analyses were performed in R 4.5.3 with lme4 1.1-37 [39], lmerTest 3.1-3, lcmm 1.9.2, brms 2.21.0 [40], JMbayes2 0.5-1 [42] (attempted then abandoned — see SAP §10 amendment 6), and nlme 3.1-164."
)

# 41 (Rizopoulos, Joint Models book) -- cite alongside the Phase III joint-model description
replace_in_para(
    "we fitted a bivariate Bayesian joint longitudinal model in brms",
    "we fitted a bivariate Bayesian joint longitudinal model in brms in which the biomarker and outcome submodels share a single random-effects covariance matrix across responses.",
    "we fitted a bivariate Bayesian joint longitudinal model [41] in brms in which the biomarker and outcome submodels share a single random-effects covariance matrix across responses."
)

# 43 (Benjamini-Hochberg FDR)
replace_in_para(
    "applied Benjamini–Hochberg FDR",
    "applied Benjamini–Hochberg FDR (q < 0.05) across the six tests, with Bonferroni (α = 0.0083) as conservative.",
    "applied Benjamini–Hochberg FDR [43] (q < 0.05) across the six tests, with Bonferroni (α = 0.0083) as conservative."
)

# 45 (Rubin, multiple imputation / Rubin's rules)
replace_in_para(
    "pooling estimates by Rubin",
    "pooling estimates by Rubin's rules.",
    "pooling estimates by Rubin's rules [45]."
)

# 44 (Vandenbroucke, STROBE explanation and elaboration)
replace_in_para(
    "This study is reported in accordance with STROBE",
    "This study is reported in accordance with STROBE for observational studies (see Supplementary File S2).",
    "This study is reported in accordance with STROBE [44] for observational studies (see Supplementary File S2)."
)

# ---------------------------------------------------------------------------
# Fix the two mediation-methodology in-text citations (added in the previous
# edit pass) to match the document's numbered-bracket convention instead of
# author-year style.
# ---------------------------------------------------------------------------
p_phase4 = next(p for p in d.paragraphs if p.text.startswith("To decompose the total APOE ε4 effect"))
full = p_phase4.text
full = full.replace(
    "we additionally conducted a counterfactual-based causal mediation analysis (Imai, Keele & Tingley, 2010) using the R mediation package (Tingley et al., 2014).",
    "we additionally conducted a counterfactual-based causal mediation analysis [46] using the R mediation package [47]."
)
for r in list(p_phase4.runs):
    r.text = ""
p_phase4.runs[0].text = full

d.save(PATH)
print("Citation insertions applied and saved.")
