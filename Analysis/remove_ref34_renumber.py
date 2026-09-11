import re
from docx import Document

PATH = "/Users/kuan-yuchu/Desktop/LOAD-BC/submit BC 5:31:26/sumit 6:3 26/Main_Text.docx"
d = Document(PATH)

# ---------------------------------------------------------------------------
# 1. Remove reference 34 (Crane et al., ADNI-Mem composite memory score) --
#    confirmed genuinely unused anywhere in the main text (the manuscript
#    does not use a composite memory outcome; only ADAS-Cog 13, MMSE, and
#    hippocampal volume are analyzed).
# ---------------------------------------------------------------------------
p34 = next(p for p in d.paragraphs if p.text.strip().startswith("34. Crane"))
p34._p.getparent().remove(p34._p)

# ---------------------------------------------------------------------------
# 2. Renumber every reference >= 35 down by 1, in the reference-list entries
#    themselves (the leading "NN. " prefix).
# ---------------------------------------------------------------------------
ref_prefix_re = re.compile(r"^(\d+)(\. )")
for p in d.paragraphs:
    m = ref_prefix_re.match(p.text.strip())
    if m:
        num = int(m.group(1))
        if num >= 35:
            newnum = num - 1
            old_prefix = f"{num}. "
            new_prefix = f"{newnum}. "
            # Only touch the very start of the run text
            r0 = p.runs[0]
            if r0.text.startswith(old_prefix):
                r0.text = new_prefix + r0.text[len(old_prefix):]
            else:
                # prefix might be split oddly across leading whitespace; fallback
                full = p.text
                assert full.startswith(old_prefix)
                new_full = new_prefix + full[len(old_prefix):]
                for r in list(p.runs):
                    r.text = ""
                p.runs[0].text = new_full

# ---------------------------------------------------------------------------
# 3. Renumber every in-text bracket citation number >= 35 down by 1.
#    Bracket contents can be a single number, a comma list, or an en-dash
#    range (e.g. "[6–8]", "[14–16,24]", "[46]"). Rebuild each bracket's
#    contents with renumbered values, preserving comma/range structure.
# ---------------------------------------------------------------------------
def renumber_token(n):
    return n - 1 if n >= 35 else n

def renumber_bracket(match):
    inner = match.group(1)
    parts = inner.split(',')
    new_parts = []
    for part in parts:
        part = part.strip()
        sep = '–' if '–' in part else ('-' if '-' in part else None)
        if sep:
            a, b = part.split(sep)
            new_parts.append(f"{renumber_token(int(a))}{sep}{renumber_token(int(b))}")
        else:
            new_parts.append(str(renumber_token(int(part))))
    return "[" + ",".join(new_parts) + "]"

bracket_re = re.compile(r"\[([0-9,\-–]+)\]")
for p in d.paragraphs:
    if not bracket_re.search(p.text):
        continue
    old_full = p.text
    new_full = bracket_re.sub(renumber_bracket, old_full)
    if new_full != old_full:
        for r in list(p.runs):
            r.text = ""
        p.runs[0].text = new_full

d.save(PATH)
print("Reference 34 removed; all subsequent numbers renumbered (list + in-text).")
