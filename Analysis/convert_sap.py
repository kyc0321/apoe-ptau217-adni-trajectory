# -*- coding: utf-8 -*-
"""Convert SAP_v1.md to a blinded, docx-format Supplementary File S1.
Fixes: (1) blinded per the journal's double-blind rule (author name/degree
removed, named collaborator genericized -- same treatment as Acknowledgements
in the main manuscript, which stays on the Title Page, not the blinded body);
(2) stale study-title line updated to match the current manuscript title
(this SAP predates the paper's retitling)."""
import re
from docx import Document
from docx.shared import Pt, Cm

FOLDER = "/Users/kuan-yuchu/Desktop/LOAD-BC-NS/submit NeurolSci"

with open(f"{FOLDER}/SAP_v1.md") as f:
    text = f.read()

# Blinding + stale-title fixes
text = text.replace(
    "**Study:** Plasma GFAP and p-tau217 Trajectories in Older Adults: APOE ε4-Specific Astrocytic Acceleration and Slope-Based Prognosis in ADNI",
    "**Study:** APOE ε4 Acts on Plasma p-tau217 Trajectory Primarily via Amyloid Burden: A Bayesian Joint Longitudinal Analysis of ADNI")
text = text.replace("**SAP author:** Kuan-Yu Chu, DDS, MHA, PhD", "**SAP author:** [blinded for peer review]")
text = text.replace("**Independent statistical reviewer:** Dr. Feng-Ping Chu (to be confirmed)",
                     "**Independent statistical reviewer:** [blinded for peer review]")
text = text.replace("| Sole author / SAP author | Kuan-Yu Chu | 2026-05-25 | _____________ |",
                     "| Sole author / SAP author | [blinded for peer review] | 2026-05-25 | _____________ |")
# Catch-all: any remaining mentions of the named collaborator elsewhere in the
# document (e.g. the amendment-history narrative in §10) -- same rationale as
# above, this person's name doesn't need to appear in a blinded review copy.
text = re.sub(r"Dr\.\s*Feng-Ping Chu", "[blinded for peer review]", text)
text = re.sub(r"Dr\.\s*Chu(?!'s)", "[blinded for peer review]", text)
text = text.replace("Dr. Chu's", "[blinded for peer review]'s")

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(10)
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2.54)

for line in text.split("\n"):
    s = line.strip()
    if not s:
        continue
    if s.startswith("### "):
        doc.add_heading(s[4:], level=3)
    elif s.startswith("## "):
        doc.add_heading(s[3:], level=2)
    elif s.startswith("# "):
        doc.add_heading(s[2:], level=1)
    elif s.startswith("|"):
        # simple table row -> plain paragraph (tables in this file are short, this is legible enough)
        p = doc.add_paragraph()
        p.add_run(s.strip("|").replace("|", "  |  "))
    else:
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.+?\*\*)', s)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                p.add_run(part[2:-2]).bold = True
            else:
                p.add_run(part)

doc.save(f"{FOLDER}/Supplementary_File_S1_SAP_R1.docx")
print("Saved. Blinding check:")
full = "\n".join(p.text for p in doc.paragraphs)
for leak in ["Kuan-Yu Chu", "Feng-Ping Chu", "kyc0321"]:
    print(f"  {leak!r}:", leak in full)
