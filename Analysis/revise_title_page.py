# -*- coding: utf-8 -*-
"""Fix Title Page.docx: rename Declarations->Statements and Declarations,
add missing Ethics approval / Consent statements, remove the STROBE/SAP
lines (relocated into the main text Methods section instead -- they aren't
in the journal's defined title-page Statements and Declarations categories),
and correct the word-count / tables-figures summary line."""
from docx import Document
import copy

FOLDER = "/Users/kuan-yuchu/Desktop/LOAD-BC-NS/submit NeurolSci"
doc = Document(f"{FOLDER}/Title Page.docx")
paras = doc.paragraphs

# 1. Rename "Declarations" heading
for p in paras:
    if p.text.strip() == "Declarations":
        for run in p.runs:
            run.text = ""
        p.runs[0].text = "Statements and Declarations"
        break

# 2. Fix word count / tables-figures summary lines
for p in paras:
    if p.text.startswith("Word count:"):
        for run in p.runs:
            run.text = ""
        p.runs[0].text = ("Word count: Abstract 250 / Main text 3,919 "
                           "(excluding references, tables, and figure legends)")
    if p.text.startswith("Tables / Figures:"):
        for run in p.runs:
            run.text = ""
        p.runs[0].text = ("Tables / Figures: 3 main (Table 1-3) / 3 main (Figure 1-3); "
                           "supplementary table/figure count TO BE FINALIZED -- see REVISION_NOTES.md "
                           "(the project's existing supplementary materials use inconsistent numbering "
                           "across two different files and need reconciliation before this line can be completed)")

# 3. Insert Ethics approval + Consent statements after Data Availability,
#    remove the Reporting-guidelines/SAP lines (relocated to Methods).
body = doc.element.body
target_idx = None
remove_paras = []
for i, p in enumerate(paras):
    if p.text.startswith("Data Availability:"):
        target_idx = i
    if p.text.startswith("Reporting guidelines:") or p.text.startswith("Pre-specified analysis plan:"):
        remove_paras.append(p)

assert target_idx is not None

data_avail_p = paras[target_idx]

def new_para_after(anchor_p, text):
    new_p = copy.deepcopy(anchor_p._p)
    anchor_p._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    para_obj = Paragraph(new_p, anchor_p._parent)
    for run in para_obj.runs:
        run.text = ""
    if para_obj.runs:
        para_obj.runs[0].text = text
    else:
        para_obj.add_run(text)
    return para_obj

ethics_text = ("Ethics Approval: This study used de-identified data from the Alzheimer's "
               "Disease Neuroimaging Initiative (ADNI). ADNI study procedures were approved "
               "by the institutional review board at each participating site, and the study "
               "was conducted in accordance with the Declaration of Helsinki. No new ethics "
               "approval was required for this secondary analysis.")
consent_text = ("Consent to Participate: Written informed consent was obtained from all ADNI "
                "participants (or their authorized representatives) at the time of enrollment, "
                "as described in the original ADNI protocol. No new consent was required for "
                "this secondary analysis of de-identified data.")

p1 = new_para_after(data_avail_p, ethics_text)
p2 = new_para_after(p1, consent_text)

for p in remove_paras:
    p._p.getparent().remove(p._p)

doc.save(f"{FOLDER}/Title Page_R1.docx")
print("Saved Title Page_R1.docx")

# verify
d2 = Document(f"{FOLDER}/Title Page_R1.docx")
for i, p in enumerate(d2.paragraphs):
    print(i, ':', repr(p.text[:90]))
