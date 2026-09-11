import re
from docx import Document

PATH = "/Users/kuan-yuchu/Desktop/LOAD-BC/submit BC 5:31:26/sumit 6:3 26/Main_Text.docx"

# ---------------------------------------------------------------------------
# Full reference table: number -> (key, full formatted author-year reference)
# Built from the current (numbered-Vancouver) reference list, converted to
# Elsevier author-year ("Vancouver - Name & Year") style. Journal names kept
# in their existing abbreviated form used throughout the manuscript.
# ---------------------------------------------------------------------------
REFS = {
1: ("GBD 2019 Dementia Forecasting Collaborators, 2022",
    "GBD 2019 Dementia Forecasting Collaborators, 2022. Estimation of the global prevalence of dementia in 2019 and forecasted prevalence in 2050: an analysis for the Global Burden of Disease Study 2019. Lancet Public Health 7(2), e105-e125. https://doi.org/10.1016/S2468-2667(21)00249-8"),
2: ("Veitch et al., 2019",
    "Veitch, D.P., Weiner, M.W., Aisen, P.S., et al., 2019. Understanding disease progression and improving Alzheimer's disease clinical trials: Recent highlights from the Alzheimer's Disease Neuroimaging Initiative. Alzheimers Dement. 15(1), 106-152. https://doi.org/10.1016/j.jalz.2018.08.005"),
3: ("Jack et al., 2018",
    "Jack, C.R. Jr, Bennett, D.A., Blennow, K., et al., 2018. NIA-AA Research Framework: Toward a biological definition of Alzheimer's disease. Alzheimers Dement. 14(4), 535-562. https://doi.org/10.1016/j.jalz.2018.02.018"),
4: ("Jack et al., 2024",
    "Jack, C.R. Jr, Andrews, J.S., Beach, T.G., et al., 2024. Revised criteria for diagnosis and staging of Alzheimer's disease: Alzheimer's Association Workgroup. Alzheimers Dement. 20(8), 5143-5169. https://doi.org/10.1002/alz.13859"),
5: ("Cicognola et al., 2021",
    "Cicognola, C., Janelidze, S., Hertze, J., et al., 2021. Plasma glial fibrillary acidic protein detects Alzheimer pathology and predicts future cognitive decline in mild cognitive impairment. Alzheimers Res. Ther. 13(1), 68. https://doi.org/10.1186/s13195-021-00804-9"),
6: ("Pereira et al., 2021",
    "Pereira, J.B., Janelidze, S., Smith, R., et al., 2021. Plasma GFAP is an early marker of amyloid-β but not tau pathology in Alzheimer's disease. Brain 144(11), 3505-3516. https://doi.org/10.1093/brain/awab223"),
7: ("Benedet et al., 2021",
    "Benedet, A.L., Milà-Alomà, M., Vrillon, A., et al., 2021. Differences between plasma and cerebrospinal fluid glial fibrillary acidic protein levels across the Alzheimer disease continuum. JAMA Neurol. 78(12), 1471-1483. https://doi.org/10.1001/jamaneurol.2021.3671"),
8: ("Chatterjee et al., 2021",
    "Chatterjee, P., Pedrini, S., Stoops, E., et al., 2021. Plasma glial fibrillary acidic protein is elevated in cognitively normal older adults at risk of Alzheimer's disease. Transl. Psychiatry 11(1), 27. https://doi.org/10.1038/s41398-020-01137-1"),
9: ("Bellaver et al., 2023",
    "Bellaver, B., Povala, G., Ferreira, P.C.L., et al., 2023. Astrocyte reactivity influences amyloid-β effects on tau pathology in preclinical Alzheimer's disease. Nat. Med. 29(7), 1775-1781. https://doi.org/10.1038/s41591-023-02380-x"),
10: ("Janelidze et al., 2020",
    "Janelidze, S., Mattsson, N., Palmqvist, S., et al., 2020. Plasma P-tau181 in Alzheimer's disease: relationship to other biomarkers, differential diagnosis, neuropathology and longitudinal progression to Alzheimer's dementia. Nat. Med. 26(3), 379-386. https://doi.org/10.1038/s41591-020-0755-1"),
11: ("Palmqvist et al., 2020",
    "Palmqvist, S., Janelidze, S., Quiroz, Y.T., et al., 2020. Discriminative accuracy of plasma phospho-tau217 for Alzheimer disease vs other neurodegenerative disorders. JAMA 324(8), 772-781. https://doi.org/10.1001/jama.2020.12134"),
12: ("Barthélemy et al., 2020",
    "Barthélemy, N.R., Horie, K., Sato, C., Bateman, R.J., 2020. Blood plasma phosphorylated-tau isoforms track CNS change in Alzheimer's disease. J. Exp. Med. 217(11), e20200861. https://doi.org/10.1084/jem.20200861"),
13: ("Ashton et al., 2021",
    "Ashton, N.J., Pascoal, T.A., Karikari, T.K., et al., 2021. Plasma p-tau231: a new biomarker for incipient Alzheimer's disease pathology. Acta Neuropathol. 141(5), 709-724. https://doi.org/10.1007/s00401-021-02275-6"),
14: ("Mattsson-Carlgren et al., 2020",
    "Mattsson-Carlgren, N., Janelidze, S., Palmqvist, S., et al., 2020. Longitudinal plasma p-tau217 is increased in early stages of Alzheimer's disease. Brain 143(11), 3234-3241. https://doi.org/10.1093/brain/awaa286"),
15: ("Ashton et al., 2024",
    "Ashton, N.J., Brum, W.S., Di Molfetta, G., et al., 2024. Diagnostic accuracy of a plasma phosphorylated tau 217 immunoassay for Alzheimer disease pathology. JAMA Neurol. 81(3), 255-263. https://doi.org/10.1001/jamaneurol.2023.5319"),
16: ("Therriault et al., 2023",
    "Therriault, J., Vermeiren, M., Servaes, S., et al., 2023. Association of phosphorylated tau biomarkers with amyloid PET vs tau PET. JAMA Neurol. 80(2), 188-199. https://doi.org/10.1001/jamaneurol.2022.4485"),
17: ("Janelidze et al., 2021",
    "Janelidze, S., Berron, D., Smith, R., et al., 2021. Associations of plasma phospho-tau217 levels with tau positron emission tomography in early Alzheimer disease. JAMA Neurol. 78(2), 149-156. https://doi.org/10.1001/jamaneurol.2020.4201"),
18: ("Cullen et al., 2021",
    "Cullen, N.C., Leuzy, A., Janelidze, S., et al., 2021. Plasma biomarkers of Alzheimer's disease improve prediction of cognitive decline in cognitively unimpaired elderly populations. Nat. Commun. 12(1), 3555. https://doi.org/10.1038/s41467-021-23746-0"),
19: ("Brum et al., 2023",
    "Brum, W.S., Cullen, N.C., Janelidze, S., et al., 2023. A two-step workflow based on plasma p-tau217 to screen for amyloid-β positivity with further confirmatory testing only in uncertain cases. Nat. Aging 3(9), 1079-1090. https://doi.org/10.1038/s43587-023-00471-5"),
20: ("Janelidze et al., 2023a",
    "Janelidze, S., Bali, D., Ashton, N.J., et al., 2023a. Head-to-head comparison of 10 plasma phospho-tau assays in prodromal Alzheimer disease. Brain 146(4), 1592-1601. https://doi.org/10.1093/brain/awac333"),
21: ("Farrer et al., 1997",
    "Farrer, L.A., Cupples, L.A., Haines, J.L., et al., 1997. Effects of age, sex, and ethnicity on the association between apolipoprotein E genotype and Alzheimer disease. JAMA 278(16), 1349-1356."),
22: ("Yamazaki et al., 2019",
    "Yamazaki, Y., Zhao, N., Caulfield, T.R., Liu, C.C., Bu, G., 2019. Apolipoprotein E and Alzheimer disease: pathobiology and targeting strategies. Nat. Rev. Neurol. 15(9), 501-518. https://doi.org/10.1038/s41582-019-0228-7"),
23: ("Serrano-Pozo et al., 2021",
    "Serrano-Pozo, A., Das, S., Hyman, B.T., 2021. APOE and Alzheimer's disease: advances in genetics, pathophysiology, and therapeutic approaches. Lancet Neurol. 20(1), 68-80. https://doi.org/10.1016/S1474-4422(20)30412-9"),
24: ("Fortea et al., 2024",
    "Fortea, J., Pegueroles, J., Alcolea, D., et al., 2024. APOE4 homozygosity represents a distinct genetic form of Alzheimer's disease. Nat. Med. 30(5), 1284-1291. https://doi.org/10.1038/s41591-024-02931-w"),
25: ("Liddelow et al., 2017",
    "Liddelow, S.A., Guttenplan, K.A., Clarke, L.E., et al., 2017. Neurotoxic reactive astrocytes are induced by activated microglia. Nature 541(7638), 481-487. https://doi.org/10.1038/nature21029"),
26: ("Habib et al., 2020",
    "Habib, N., McCabe, C., Medina, S., et al., 2020. Disease-associated astrocytes in Alzheimer's disease and aging. Nat. Neurosci. 23(6), 701-706. https://doi.org/10.1038/s41593-020-0624-8"),
27: ("Petersen et al., 2010",
    "Petersen, R.C., Aisen, P.S., Beckett, L.A., et al., 2010. Alzheimer's Disease Neuroimaging Initiative (ADNI): clinical characterization. Neurology 74(3), 201-209."),
28: ("Weiner et al., 2017",
    "Weiner, M.W., Veitch, D.P., Aisen, P.S., et al., 2017. The Alzheimer's Disease Neuroimaging Initiative 3: continued innovation for clinical trial improvement. Alzheimers Dement. 13(5), 561-571. https://doi.org/10.1016/j.jalz.2016.10.006"),
29: ("Trojanowski et al., 2010",
    "Trojanowski, J.Q., Vandeerstichele, H., Korecka, M., et al., 2010. Update on the biomarker core of the Alzheimer's Disease Neuroimaging Initiative subjects. Alzheimers Dement. 6(3), 230-238."),
30: ("Rissin et al., 2010",
    "Rissin, D.M., Kan, C.W., Campbell, T.G., et al., 2010. Single-molecule enzyme-linked immunosorbent assay detects serum proteins at subfemtomolar concentrations. Nat. Biotechnol. 28(6), 595-599. https://doi.org/10.1038/nbt.1641"),
31: ("Mohs et al., 1997",
    "Mohs, R.C., Knopman, D., Petersen, R.C., et al., 1997. Development of cognitive instruments for use in clinical trials of antidementia drugs: additions to the Alzheimer's Disease Assessment Scale that broaden its scope. Alzheimer Dis. Assoc. Disord. 11(Suppl 2), S13-S21."),
32: ("Folstein et al., 1975",
    "Folstein, M.F., Folstein, S.E., McHugh, P.R., 1975. “Mini-mental state”: A practical method for grading the cognitive state of patients for the clinician. J. Psychiatr. Res. 12(3), 189-198."),
33: ("Fischl, 2012",
    "Fischl, B., 2012. FreeSurfer. Neuroimage 62(2), 774-781. https://doi.org/10.1016/j.neuroimage.2012.01.021"),
34: ("Mielke et al., 2022",
    "Mielke, M.M., Dage, J.L., Frank, R.D., et al., 2022. Performance of plasma phosphorylated tau 181 and 217 in the community. Nat. Med. 28(7), 1398-1405. https://doi.org/10.1038/s41591-022-01822-2"),
35: ("Inker et al., 2021",
    "Inker, L.A., Eneanya, N.D., Coresh, J., et al., 2021. New creatinine- and cystatin C-based equations to estimate GFR without race. N. Engl. J. Med. 385(19), 1737-1749. https://doi.org/10.1056/NEJMoa2102953"),
36: ("Janelidze et al., 2023b",
    "Janelidze, S., Barthélemy, N.R., He, Y., Bateman, R.J., Hansson, O., 2023b. Mitigating the associations of kidney dysfunction with blood biomarkers of Alzheimer disease by using phosphorylated tau to total tau ratios. JAMA Neurol. 80(5), 516-522. https://doi.org/10.1001/jamaneurol.2023.0199"),
37: ("Proust-Lima et al., 2017",
    "Proust-Lima, C., Philipps, V., Liquet, B., 2017. Estimation of extended mixed models using latent classes and latent processes: the R package lcmm. J. Stat. Softw. 78(2), 1-56. https://doi.org/10.18637/jss.v078.i02"),
38: ("Bates et al., 2015",
    "Bates, D., Mächler, M., Bolker, B., Walker, S., 2015. Fitting linear mixed-effects models using lme4. J. Stat. Softw. 67(1), 1-48. https://doi.org/10.18637/jss.v067.i01"),
39: ("Bürkner, 2017",
    "Bürkner, P.C., 2017. brms: An R package for Bayesian multilevel models using Stan. J. Stat. Softw. 80(1), 1-28. https://doi.org/10.18637/jss.v080.i01"),
40: ("Rizopoulos, 2012",
    "Rizopoulos, D., 2012. Joint Models for Longitudinal and Time-to-Event Data: With Applications in R. Chapman & Hall/CRC, Boca Raton, FL."),
41: ("Rizopoulos et al., 2024",
    "Rizopoulos, D., Papageorgiou, G., Miranda Afonso, P., 2024. JMbayes2: Extended Joint Models for Longitudinal and Time-to-Event Data. R package v0.5-1."),
42: ("Benjamini and Hochberg, 1995",
    "Benjamini, Y., Hochberg, Y., 1995. Controlling the false discovery rate: a practical and powerful approach to multiple testing. J. R. Stat. Soc. B 57(1), 289-300."),
43: ("Vandenbroucke et al., 2007",
    "Vandenbroucke, J.P., von Elm, E., Altman, D.G., et al., 2007. STROBE: explanation and elaboration. PLoS Med. 4(10), e297."),
44: ("Rubin, 1987",
    "Rubin, D.B., 1987. Multiple Imputation for Nonresponse in Surveys. John Wiley & Sons, New York."),
45: ("Imai et al., 2010",
    "Imai, K., Keele, L., Tingley, D., 2010. A general approach to causal mediation analysis. Psychol. Methods 15(4), 309-334. https://doi.org/10.1037/a0020761"),
46: ("Tingley et al., 2014",
    "Tingley, D., Yamamoto, T., Hirose, K., Keele, L., Imai, K., 2014. mediation: R package for causal mediation analysis. J. Stat. Softw. 59(5), 1-38. https://doi.org/10.18637/jss.v059.i05"),
47: ("Woo et al., 2025",
    "Woo, M.S., Macedo, A.C., Hosseini, S.A., et al., 2025. APOE ε4 carriership determines a faster plasma p-tau217 progression in Aβ-positive individuals. Alzheimers Dement. 21(12), e71048. https://doi.org/10.1002/alz.71048"),
}

def sort_key_for_alpha(entry):
    key, _ = entry
    # sort alphabetically by the key text, case-insensitive
    return key.lower()

def bracket_to_authoryear(match):
    inner = match.group(1)
    nums = []
    for part in inner.split(','):
        part = part.strip()
        sep = '–' if '–' in part else ('-' if '-' in part else None)
        if sep:
            a, b = map(int, part.split(sep))
            nums.extend(range(a, b + 1))
        elif part:
            nums.append(int(part))
    # de-duplicate, preserve ascending numeric order (approximates original intent)
    seen = []
    for n in nums:
        if n not in seen:
            seen.append(n)
    keys = [REFS[n][0] for n in seen]
    return "(" + "; ".join(keys) + ")"

d = Document(PATH)
bracket_re = re.compile(r"\[([0-9,\-–]+)\]")

changed = 0
for p in d.paragraphs:
    if not bracket_re.search(p.text):
        continue
    old_full = p.text
    new_full = bracket_re.sub(bracket_to_authoryear, old_full)
    if new_full != old_full:
        for r in list(p.runs):
            r.text = ""
        p.runs[0].text = new_full
        changed += 1

print(f"Converted in-text citations in {changed} paragraph(s).")

# ---------------------------------------------------------------------------
# Rebuild the References section: remove the old numbered paragraphs, insert
# new alphabetically-sorted author-year entries in their place.
# ---------------------------------------------------------------------------
paras = list(d.paragraphs)
ref_heading = next(p for p in paras if p.text.strip() == "References")
# find all paragraphs after the heading that are numbered reference entries
old_ref_paras = []
started = False
for p in paras:
    if p is ref_heading:
        started = True
        continue
    if started:
        if re.match(r"^\d+\. ", p.text.strip()):
            old_ref_paras.append(p)

assert len(old_ref_paras) == 47, f"expected 47 old reference paragraphs, found {len(old_ref_paras)}"

sorted_entries = sorted(REFS.values(), key=sort_key_for_alpha)

# Insert new entries right before the first old ref paragraph, in sorted order
anchor = old_ref_paras[0]
for key, full_ref in sorted_entries:
    new_p = anchor.insert_paragraph_before("")
    new_p.add_run(full_ref)

# Now remove all the old numbered paragraphs
for p in old_ref_paras:
    p._p.getparent().remove(p._p)

d.save(PATH)
print("References section rebuilt in alphabetical author-year order.")
